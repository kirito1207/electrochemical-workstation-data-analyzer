from __future__ import annotations

import csv

from openpyxl import load_workbook

from export.csv import export_csv_bundle
from export.excel import export_analysis_workbook
from export.logging import export_analysis_settings
from plotting.selected_potential import plot_selected_potential


def test_csv_exports_include_traceable_rows_and_parsed_curves(tmp_path, synthetic_analysis_result):
    files = export_csv_bundle(synthetic_analysis_result, tmp_path / "csv")
    selected = tmp_path / "csv" / "selected_potential_data.csv"
    with selected.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))

    assert len(rows) == 42
    assert all(row["source_sha256"] for row in rows)
    assert all("signed_current_A" in row and "response_magnitude_uA" in row for row in rows)
    assert len(list((tmp_path / "csv" / "parsed_lsv").glob("*.csv"))) == 42
    assert len(files) == 49


def test_excel_contains_all_required_sheets(tmp_path, synthetic_analysis_result):
    path = export_analysis_workbook(synthetic_analysis_result, tmp_path / "LSV_analysis.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=False)

    assert workbook.sheetnames == [
        "Analysis_Settings",
        "File_Metadata",
        "Experiment_Parameters",
        "Selected_Potential_Data",
        "Group_Summary",
        "Statistics",
        "Outlier_Flags",
        "Exclusion_Log",
    ]
    assert workbook["Selected_Potential_Data"].max_row == 43
    assert workbook["Exclusion_Log"]["A2"].value == "All data included"


def test_analysis_log_records_settings_and_all_sources(tmp_path, synthetic_analysis_result):
    path = export_analysis_settings(synthetic_analysis_result, tmp_path / "analysis_settings.json")
    text = path.read_text(encoding="utf-8")

    assert '"target_potential_V": 0.0' in text
    assert '"analysis_metric": "magnitude"' in text
    assert text.count('"source_sha256"') == 42


def test_selected_potential_figures_export_png_svg_pdf(tmp_path, synthetic_analysis_result):
    files = plot_selected_potential(synthetic_analysis_result, tmp_path / "figures")

    assert len(files) == 6
    assert {path.suffix for path in files} == {".png", ".svg", ".pdf"}
    assert all(path.stat().st_size > 0 for path in files)
