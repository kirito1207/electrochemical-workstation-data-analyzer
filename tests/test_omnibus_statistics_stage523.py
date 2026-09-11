from __future__ import annotations

from dataclasses import replace
from inspect import getsource
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest
from scipy import stats

from analysis import (
    AnalysisSettings,
    ComparisonDefinition,
    CurrentAtPotential,
    ManifestEntry,
    analyze_lsv_with_manifest,
    confirmed_generic_manifest,
    kruskal_wallis,
    welch_anova,
)
from chi_gui.lsv_workflow import (
    LSVWorkflowState,
    StaleAnalysisResultError,
    omnibus_display_rows,
    omnibus_result_status,
)
from chi_gui.widgets.lsv_analysis import LSVResultsPanel, LSVSettingsPanel
from chi_gui.workspaces import WorkspaceManager
from export.csv import export_csv_bundle
from export.excel import export_analysis_workbook


REFERENCE_GROUPS = {
    "A": [1.0, 2.0, 3.0, 4.0],
    "B": [2.0, 4.0, 6.0, 8.0, 10.0],
    "C": [3.0, 3.0, 4.0, 5.0, 9.0, 11.0],
}


def test_welch_anova_known_reference_case():
    result = welch_anova(REFERENCE_GROUPS)
    assert result.status == "ok"
    assert result.statistic == pytest.approx(3.8557896342688545)
    assert result.df1 == 2.0
    assert result.df2 == pytest.approx(7.50107874926156)
    assert result.p_value == pytest.approx(0.07051462528897864)


def test_welch_anova_matches_scipy_for_unequal_variances():
    groups = {"A": [1, 2, 3], "B": [2, 8, 14, 20], "C": [4, 4.5, 5, 5.5, 6]}
    result = welch_anova(groups)
    assert result.statistic == pytest.approx(9.846533777433473)
    assert result.p_value == pytest.approx(0.02435844382588857)


def test_welch_anova_handles_extreme_variance_differences():
    groups = {
        "tight": [1.0, 1.0001, 0.9999, 1.0002],
        "wide": [-1000.0, -1.0, 2.0, 1200.0],
        "middle": [2.0, 4.0, 6.0, 8.0],
    }
    result = welch_anova(groups)
    assert result.status == "ok"
    assert result.statistic == pytest.approx(4.1193165053812155)
    assert result.p_value == pytest.approx(0.10682038953528225)


@pytest.mark.parametrize("count", (3, 4))
def test_welch_anova_supports_three_and_four_groups(count):
    groups = {f"G{i}": np.arange(i, i + 5, dtype=float) for i in range(count)}
    result = welch_anova(groups)
    assert result.status == "ok"
    assert result.group_count == count
    assert result.total_n == count * 5


def test_omnibus_n_less_than_two_is_structured_unavailable():
    groups = {"A": [1, 2], "B": [3], "C": [4, 5]}
    for result in (welch_anova(groups), kruskal_wallis(groups)):
        assert result.status == "unavailable"
        assert "B (n=1)" in result.notes


def test_welch_zero_variance_is_structured_unavailable():
    result = welch_anova({"A": [1, 1], "B": [2, 3], "C": [4, 6]})
    assert result.status == "unavailable"
    assert "zero within-group variance" in result.notes


def test_identical_groups_do_not_return_silent_nan():
    groups = {"A": [1, 1], "B": [1, 1], "C": [1, 1]}
    assert welch_anova(groups).status == "unavailable"
    kw = kruskal_wallis(groups)
    assert kw.status == "unavailable"
    assert kw.statistic is None and kw.p_value is None


def test_kruskal_wallis_known_reference_case():
    result = kruskal_wallis(REFERENCE_GROUPS)
    assert result.status == "ok"
    assert result.statistic == pytest.approx(4.165910465819716)
    assert result.df1 == 2.0 and result.df2 is None
    assert result.p_value == pytest.approx(0.12456155931991668)


@pytest.mark.parametrize("count", (3, 4))
def test_kruskal_wallis_supports_three_and_four_groups_with_ties(count):
    groups = {f"G{i}": [i, i + 1, i + 1, i + 2] for i in range(count)}
    expected = stats.kruskal(*groups.values())
    result = kruskal_wallis(groups)
    assert result.group_count == count
    assert result.statistic == pytest.approx(expected.statistic)
    assert result.p_value == pytest.approx(expected.pvalue)


def test_nonfinite_values_are_reported_without_dropping_a_group():
    groups = {"A": [1, 2], "B": [3, np.nan], "C": [4, 5]}
    for result in (welch_anova(groups), kruskal_wallis(groups)):
        assert result.status == "unavailable"
        assert result.group_count == 3
        assert "B" in result.notes


def _analysis_result(monkeypatch, grouped_signed, *, metric="magnitude", bare=()):
    values_by_path = {}
    entries = []
    for group, values in grouped_signed.items():
        for index, value in enumerate(values, start=1):
            path = Path(f"{group}-S{index}.bin")
            values_by_path[str(path)] = float(value)
            entries.append(
                ManifestEntry(path.name, str(path), group, "Material", f"S{index}", file_path=str(path))
            )
    for index, (group, value) in enumerate(bare, start=1):
        path = Path(f"{group}-Bare{index}.bin")
        values_by_path[str(path)] = float(value)
        entries.append(
            ManifestEntry(path.name, str(path), group, "Bare", f"Bare{index}", file_path=str(path))
        )

    def fake_parse(path):
        return SimpleNamespace(
            file_name=Path(path).name,
            file_path=Path(path),
            source_sha256="synthetic",
            experiment_type="Linear Sweep Voltammetry",
            file_size_bytes=8,
            n_points=2,
            configured_start_potential_V=-0.1,
            configured_final_potential_V=0.1,
            actual_first_potential_V=-0.1,
            actual_last_potential_V=0.1,
            scan_rate_V_s=0.1,
            potential_increment_V=0.2,
            potential_V=np.asarray([-0.1, 0.1], dtype=float),
            current_A=np.asarray([0.0, 0.0], dtype=float),
            data_start_byte=0,
            current_encoding="synthetic float64",
            validation_status="valid",
        )

    def fake_extract(data, target):
        value = values_by_path[data.file_name]
        return CurrentAtPotential(
            data.file_name, data.file_name, target, value * 1e-6, value, abs(value),
            False, target, target,
        )

    monkeypatch.setattr("analysis.lsv_analysis.parse_lsv", fake_parse)
    monkeypatch.setattr("analysis.lsv_analysis.extract_current_at_potential", fake_extract)
    manifest = confirmed_generic_manifest(entries)
    return analyze_lsv_with_manifest(
        manifest,
        settings=AnalysisSettings(analysis_metric=metric, analysis_timestamp="2026-09-11T00:00:00+00:00"),
    )


@pytest.mark.parametrize("metric", ("signed", "magnitude"))
def test_omnibus_uses_the_selected_metric(monkeypatch, metric):
    signed = {"A": [-4, -2, -1], "B": [-2, 1, 3], "C": [1, 4, 8]}
    result = _analysis_result(monkeypatch, signed, metric=metric)
    expected_values = signed if metric == "signed" else {
        group: [abs(value) for value in values] for group, values in signed.items()
    }
    expected = welch_anova(expected_values)
    actual = next(item for item in result.omnibus_tests if item.test == "Welch ANOVA")
    assert actual.statistic == pytest.approx(expected.statistic)
    assert actual.p_value == pytest.approx(expected.p_value)


def test_all_material_groups_are_included_and_bare_is_excluded(monkeypatch):
    groups = {"A": [1, 2], "B": [3, 5], "C": [6, 9], "D": [10, 14]}
    result = _analysis_result(monkeypatch, groups, bare=(("A", 1000), ("D", -1000)))
    assert result.groups == ("A", "B", "C", "D")
    assert all(item.group_count == 4 and item.total_n == 8 for item in result.omnibus_tests)


def test_omnibus_p_does_not_gate_explicit_pairwise(monkeypatch):
    groups = {"A": [1, 2, 3], "B": [1.1, 2.1, 3.1], "C": [0.9, 1.9, 2.9]}
    monkeypatch.setattr(
        "analysis.statistics._bootstrap_intervals", lambda *args, **kwargs: (0, 0, 0, 0)
    )
    base = _analysis_result(monkeypatch, groups)
    result = analyze_lsv_with_manifest(
        base.manifest,
        comparisons=(ComparisonDefinition("A", "B", "primary", "primary", "A-B"),),
        settings=base.settings,
    )
    assert next(item for item in result.omnibus_tests if item.test == "Welch ANOVA").p_value > 0.05
    assert {item.comparison for item in result.comparisons} == {"A-B"}
    assert next(item for item in result.comparisons if item.test.startswith("Welch")).holm_adjusted_p is not None


@pytest.mark.parametrize("group_count", (1, 2, 3, 4, 7))
def test_gui_overall_status_and_rows_support_one_two_three_four_and_n(monkeypatch, group_count):
    groups = {f"G{i + 1}": [i + 1, i + 2] for i in range(group_count)}
    result = _analysis_result(monkeypatch, groups)
    text = omnibus_result_status(result)
    rows = omnibus_display_rows(result)
    assert len(rows) == 2
    if group_count < 3:
        assert "不适用" in text or "需 ≥3" in text
        assert {row["status"] for row in rows} == {"不适用"}
    else:
        assert f"Groups：{group_count}" in text
        assert {row["test"] for row in rows} == {"Welch ANOVA", "Kruskal-Wallis"}


def test_overall_results_are_additive_and_pairwise_editor_is_retained():
    results_source = getsource(LSVResultsPanel.__init__)
    settings_source = getsource(LSVSettingsPanel.__init__)
    assert '"整体多组比较"' in results_source
    assert '"用户定义组间比较"' in results_source
    for label in ("self.left", "self.right", "self.role", "self.family", "self.comparison_name"):
        assert label in settings_source
    assert 'text="用户定义组间比较"' in settings_source
    assert '"Holm Family"' in settings_source


def test_omnibus_export_adds_csv_and_excel_sheet(monkeypatch, tmp_path):
    result = _analysis_result(monkeypatch, {"A": [1, 2], "B": [3, 4], "C": [5, 7]})
    generated = export_csv_bundle(result, tmp_path / "csv")
    csv_path = tmp_path / "csv" / "omnibus_statistics.csv"
    assert csv_path in generated
    import csv

    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert [row["test"] for row in rows] == ["Welch ANOVA", "Kruskal-Wallis"]
    assert all(row["analysis_metric"] == "magnitude" for row in rows)
    assert all(row["group_count"] == "3" and row["total_n"] == "6" for row in rows)
    assert all(row["status"] == "ok" for row in rows)
    workbook_path = export_analysis_workbook(result, tmp_path / "LSV_analysis.xlsx")
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True)
    assert "Omnibus statistics" in workbook.sheetnames
    assert workbook["Omnibus statistics"].max_row == 3


def test_single_group_export_allows_zero_pairwise_rows(monkeypatch, tmp_path):
    result = _analysis_result(monkeypatch, {"Only": [1, 2]})
    export_csv_bundle(result, tmp_path)
    assert (tmp_path / "statistics.csv").read_text(encoding="utf-8-sig").count("\n") == 1


def test_omnibus_result_obeys_existing_stale_export_policy(monkeypatch):
    result = _analysis_result(monkeypatch, {"A": [1, 2], "B": [3, 4], "C": [5, 7]})
    workflow = LSVWorkflowState(
        analysis_result=result,
        result_signature=("completed",),
    )
    workflow.set_metric("signed")
    assert workflow.result_stale
    with pytest.raises(StaleAnalysisResultError):
        workflow.require_exportable_result()


def test_omnibus_results_remain_workspace_local(monkeypatch):
    four_group = _analysis_result(
        monkeypatch, {"A": [1, 2], "B": [3, 4], "C": [5, 7], "D": [8, 11]}
    )
    two_group = _analysis_result(monkeypatch, {"X": [1, 2], "Y": [3, 4]})
    manager = WorkspaceManager()
    first = manager.active
    first.lsv_workflow.analysis_result = four_group
    second = manager.create()
    second.lsv_workflow.analysis_result = two_group
    assert first.lsv_workflow.analysis_result.omnibus_tests[0].group_count == 4
    assert second.lsv_workflow.analysis_result.omnibus_tests[0].status == "not_applicable"
