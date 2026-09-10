from __future__ import annotations

import csv
import json

from openpyxl import load_workbook

from export import export_it_analysis_log, export_it_csv_bundle, export_it_workbook
from plotting import plot_annotated_it, plot_it_calibrations, plot_raw_it


EXPECTED_SHEETS = [
    "Analysis_Settings",
    "File_Metadata",
    "Step_Protocol",
    "Interval_Definitions",
    "Plateau_Results",
    "DeltaI_Data",
    "Calibration",
    "Concentration_Summary",
    "Outlier_Flags",
    "Current_Direction_QC",
    "Exclusion_Log",
]


def test_it_csv_outputs_include_raw_protocol_plateau_delta_and_calibration(
    tmp_path, synthetic_it_batch_result
):
    files = export_it_csv_bundle(synthetic_it_batch_result, tmp_path / "csv")

    assert len(list((tmp_path / "csv" / "raw_it").glob("*.csv"))) == 3
    assert {path.name for path in files}.issuperset(
        {
            "step_protocol.csv",
            "plateau_results.csv",
            "deltaI_data.csv",
            "calibration.csv",
            "concentration_summary.csv",
            "outlier_flags.csv",
        }
    )
    with (tmp_path / "csv" / "deltaI_data.csv").open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 12
    assert "signed_delta_I_uA" in rows[0]
    assert "magnitude_delta_I_uA" in rows[0]


def test_it_excel_contains_all_required_sheets(tmp_path, synthetic_it_batch_result):
    path = export_it_workbook(synthetic_it_batch_result, tmp_path / "IT_analysis.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=False)

    assert workbook.sheetnames == EXPECTED_SHEETS
    assert workbook["Plateau_Results"].max_row == 13
    assert workbook["Exclusion_Log"]["A2"].value == "All data included"
    assert workbook["Analysis_Settings"]["B12"].value.startswith("LOD not calculated")


def test_it_log_records_protocol_windows_qc_and_provenance(tmp_path, synthetic_it_batch_result):
    path = export_it_analysis_log(synthetic_it_batch_result, tmp_path / "analysis.json")
    payload = json.loads(path.read_text(encoding="utf-8"))

    assert payload["settings"]["plateau_fraction"] == 0.2
    assert payload["lod"]["status"] == "LOD not calculated"
    assert len(payload["files"]) == 3
    assert len(payload["files"][0]["actual_plateau_windows"]) == 4
    assert payload["files"][0]["source_sha256"]
    assert "current_direction_qc" in payload


def test_it_figures_export_png_svg_pdf(tmp_path, synthetic_it_batch_result):
    raw = plot_raw_it(synthetic_it_batch_result, tmp_path / "raw")
    annotated = plot_annotated_it(synthetic_it_batch_result, tmp_path / "annotated")
    calibration = plot_it_calibrations(synthetic_it_batch_result, tmp_path / "calibration")

    assert len(raw) == 9
    assert len(annotated) == 9
    assert len(calibration) == 39
    assert {path.suffix for path in (*raw, *annotated, *calibration)} == {".png", ".svg", ".pdf"}
    assert all(path.stat().st_size > 0 for path in (*raw, *annotated, *calibration))
