"""Formatted Excel workbook for traceable LSV analysis outputs."""

from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis.lsv_analysis import LSVAnalysisResult

from .csv import (
    STATISTICS_FIELDS,
    _omnibus_rows,
    _selected_rows,
    _statistics_rows,
    _summary_rows,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="222222")
THIN_BLUE = Side(style="thin", color="B4C7E7")

EXPLICIT_WIDTHS = {
    "file_name": 36,
    "source_file": 36,
    "sample_id": 13,
    "relative_path": 52,
    "source_sha256": 66,
    "current_encoding": 38,
    "comparison_role": 40,
    "test": 46,
    "multiple_comparison": 35,
    "outlier_method": 25,
    "outlier_reason": 58,
    "warning": 72,
    "warnings": 58,
    "notes": 32,
    "protocol_source": 34,
    "included_concentrations_uM": 30,
    "lod_reason": 58,
    "error": 58,
    "statistical_methods": 58,
    "Value": 58,
}
INTEGER_FIELDS = {
    "file_size_bytes",
    "n_points",
    "data_start_byte",
    "n",
    "bootstrap_seed",
    "bootstrap_resamples",
    "negative_count",
    "positive_count",
    "near_zero_count",
    "total_count",
    "included_count",
    "excluded_count",
    "plateau_n_points",
}


def _number_format(header: str) -> str | None:
    if header in INTEGER_FIELDS:
        return "0"
    if header.endswith("_A"):
        return "0.000000E+00"
    if any(
        token in header
        for token in (
            "_V",
            "_uA",
            "_uM",
            "_s",
            "mean",
            "median",
            "sd",
            "sem",
            "cv_percent",
            "minimum",
            "maximum",
            "q1",
            "q3",
            "iqr",
            "statistic",
            "raw_p",
            "holm_adjusted_p",
            "effect_size",
            "ci_",
            "response",
            "slope",
            "intercept",
            "r_squared",
            "cv_",
        )
    ):
        return "0.000000"
    return None


def _append_table(sheet, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    sheet.append(tuple(headers))
    for row in rows:
        sheet.append(tuple(row))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_BLUE)
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for column, cell in enumerate(row, start=1):
            cell.font = BODY_FONT
            header = str(headers[column - 1])
            wrap = header in {
                "relative_path",
                "comparison_role",
                "test",
                "multiple_comparison",
                "outlier_reason",
                "warning",
                "warnings",
                "notes",
                "protocol_source",
                "lod_reason",
                "error",
                "Value",
            }
            cell.alignment = Alignment(vertical="center", wrap_text=wrap)
            number_format = _number_format(header)
            if number_format and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = number_format
        if any(
            str(headers[column - 1]) in {
                "comparison_role",
                "test",
                "outlier_reason",
                "warning",
                "warnings",
                "notes",
                "protocol_source",
                "lod_reason",
                "error",
                "Value",
            }
            for column in range(1, len(headers) + 1)
        ):
            sheet.row_dimensions[row[0].row].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for column in range(1, sheet.max_column + 1):
        header = str(headers[column - 1])
        values = [
            "" if sheet.cell(row, column).value is None else str(sheet.cell(row, column).value)
            for row in range(1, min(sheet.max_row, 80) + 1)
        ]
        width = EXPLICIT_WIDTHS.get(
            header,
            min(max(max(map(len, values), default=8) + 2, 11), 24),
        )
        sheet.column_dimensions[get_column_letter(column)].width = width


def _dict_table(
    sheet, rows: list[dict[str, object]], headers: Sequence[str] = ()
) -> None:
    headers = tuple(headers) or (tuple(rows[0]) if rows else ())
    _append_table(sheet, headers, ([row[key] for key in headers] for row in rows))


def export_analysis_workbook(result: LSVAnalysisResult, path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    settings_sheet = workbook.create_sheet("Analysis_Settings")
    setting_rows = [
        ("target_potential_V", result.settings.target_potential_V),
        ("analysis_metric", result.settings.analysis_metric),
        ("analysis_timestamp", result.settings.analysis_timestamp),
        ("primary_test", "Welch independent-samples t-test"),
        ("sensitivity_test", "Mann-Whitney U, two-sided"),
        ("overall_primary_test", "One-way Welch ANOVA using all Material Groups"),
        ("overall_sensitivity_test", "Kruskal-Wallis using all Material Groups"),
        (
            "multiple_comparison",
            "Holm adjustment only within explicitly declared primary families",
        ),
        ("manifest_user_confirmed", result.manifest.user_confirmed),
        ("manifest_source", result.manifest.source),
        ("experiment_template", result.manifest.template_name or "Generic LSV Mode"),
        ("effect_size", "Hedges' g; percentile bootstrap 95% CI"),
        ("bootstrap_seed", result.settings.bootstrap_seed),
        ("bootstrap_resamples", result.settings.bootstrap_resamples),
        ("outlier_method", result.settings.outlier_method),
        ("sign_zero_tolerance_A", result.settings.sign_zero_tolerance_A),
        ("sign_qc_warnings", " | ".join(result.warnings) if result.warnings else "None"),
        ("exclusions", "All data included"),
        ("software_version", result.settings.software_version),
    ]
    _append_table(settings_sheet, ("Setting", "Value"), setting_rows)

    sign_qc_rows = [asdict(item) for item in result.current_sign_qc]
    sign_qc_sheet = workbook.create_sheet("Current_Sign_QC")
    _dict_table(sign_qc_sheet, sign_qc_rows)

    metadata_rows = [
        {
            **asdict(item.manifest),
            "source_sha256": item.data.source_sha256,
            "validation_status": item.data.validation_status,
        }
        for item in result.files
    ]
    _dict_table(workbook.create_sheet("File_Metadata"), metadata_rows)

    parameter_rows = [
        {
            "file_name": item.manifest.file_name,
            "relative_path": item.manifest.relative_path,
            "experiment_type": item.data.experiment_type,
            "file_size_bytes": item.data.file_size_bytes,
            "n_points": item.data.n_points,
            "configured_start_potential_V": item.data.configured_start_potential_V,
            "configured_final_potential_V": item.data.configured_final_potential_V,
            "actual_first_potential_V": item.data.actual_first_potential_V,
            "actual_last_potential_V": item.data.actual_last_potential_V,
            "scan_rate_V_s": item.data.scan_rate_V_s,
            "potential_increment_V": item.data.potential_increment_V,
            "data_start_byte": item.data.data_start_byte,
            "current_encoding": item.data.current_encoding,
        }
        for item in result.files
    ]
    _dict_table(workbook.create_sheet("Experiment_Parameters"), parameter_rows)
    _dict_table(workbook.create_sheet("Selected_Potential_Data"), _selected_rows(result))
    _dict_table(workbook.create_sheet("Group_Summary"), _summary_rows(result))
    _dict_table(workbook.create_sheet("Omnibus statistics"), _omnibus_rows(result))
    _dict_table(
        workbook.create_sheet("Statistics"),
        _statistics_rows(result),
        STATISTICS_FIELDS,
    )

    outlier_rows = [asdict(item) for item in result.outlier_flags]
    outlier_headers = (
        "file_name",
        "group",
        "sample_id",
        "target_potential_V",
        "response",
        "analysis_metric",
        "outlier_method",
        "outlier_reason",
        "status",
    )
    outlier_sheet = workbook.create_sheet("Outlier_Flags")
    _append_table(
        outlier_sheet,
        outlier_headers,
        ([row[key] for key in outlier_headers] for row in outlier_rows),
    )
    exclusion_sheet = workbook.create_sheet("Exclusion_Log")
    _append_table(
        exclusion_sheet,
        ("status", "file_name", "reason"),
        (("All data included", "", ""),),
    )

    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.tabColor = "1F4E78" if sheet.title in {
            "Analysis_Settings", "Current_Sign_QC", "Group_Summary",
            "Omnibus statistics", "Statistics"
        } else "9DC3E6"

    workbook.save(output)
    return output
